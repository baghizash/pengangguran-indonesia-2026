#!/usr/bin/env python3
"""Build Portofolio-Analisis-Pengangguran-2026.pptx (16:9, 10 slides, ID).

Redesign design-taste: Tema T2 (oat, blue, pink), varian Soft Premium
(double-bezel, sudut rounded, pill stat). Semua angka dari findings.json,
identik dengan deck lama; yang berubah hanya sistem visual.
"""
import json
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
F = json.loads((HERE / "findings.json").read_text())
CHARTS = HERE / "charts"
OUT = Path.home() / "workspace" / "your_files" / "Portofolio-Analisis-Pengangguran-2026.pptx"
OUT_COPY = HERE / "Portofolio-Analisis-Pengangguran-2026.pptx"

# ---- Tema T2: OAT / BLUE / PINK ----
GROUND = RGBColor(0xEF, 0xE7, 0xD8)    # oat
SURFACE = RGBColor(0xF7, 0xF1, 0xE5)   # permukaan terang
INK = RGBColor(0x2E, 0x31, 0x38)
BLUE = RGBColor(0x4E, 0x7F, 0xA6)      # primary
PINK = RGBColor(0xC7, 0x5B, 0x7F)      # aksen
PINKSOFT = RGBColor(0xEF, 0xC3, 0xD2)  # pink lembut
MUTED = RGBColor(0x5F, 0x5B, 0x4F)     # teks sekunder hangat
LINEC = RGBColor(0xDE, 0xD2, 0xBA)     # garis panel
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Noto Sans"

ASPECT = {  # lebar/tinggi PNG chart hasil regenerasi
    "01_tpt_nasional_tren": 1.839, "02_paradoks_pendidikan": 1.841,
    "03_pendidikan_tren": 1.839, "04_umur": 1.839,
    "05_kesenjangan": 2.203, "06_provinsi_top10": 1.678,
    "09_formal_informal": 2.089, "10_upah_sektor": 2.405,
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def jt(v):
    return f"{v} jt".replace(".", ",")


def pct(v):
    return f"{v}%".replace(".", ",")


def num(v):
    return f"{v}".replace(".", ",")


# ---------- primitives ----------
def new_slide():
    s = prs.slides.add_slide(blank)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = GROUND
    return s


def rrect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.16):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    except Exception:
        pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def bezel(slide, x, y, w, h):
    """Soft Premium double-bezel: panel permukaan + garis dalam pink lembut."""
    rrect(slide, x, y, w, h, fill=SURFACE, line=LINEC, line_w=1.0, radius=0.18)
    rrect(slide, x + 0.09, y + 0.09, w - 0.18, h - 0.18,
          fill=None, line=PINKSOFT, line_w=1.2, radius=0.16)


def pill(slide, x, y, w, h, text, fill, tcolor, size=12.5):
    rrect(slide, x, y, w, h, fill=fill, radius=h / 2)
    tf = textbox(slide, x + 0.05, y, w - 0.1, h)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_par(tf, text, size, tcolor, bold=True, first=True, align=PP_ALIGN.CENTER)
    tf.paragraphs[0].space_after = Pt(0)
    return tf


def textbox(slide, l, t, w, h):
    tx = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def add_par(tf, text, size, color, bold=False, italic=False, first=False, align=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(8)
    p.line_spacing = 1.07
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = FONT
    r.font.color.rgb = color
    return p


def add_bullet(tf, text, size=16, first=False):
    p = add_par(tf, text, size, INK, first=first)
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))
    pPr.set("marL", "171450")
    pPr.set("indent", "-171450")
    p.space_after = Pt(11)
    return p


def bullets(slide, x, y, w, h, items, size=16):
    tf = textbox(slide, x, y, w, h)
    for i, b in enumerate(items):
        add_bullet(tf, b, size=size, first=(i == 0))
    return tf


def title_block(slide, text):
    tf = textbox(slide, 0.55, 0.30, 12.2, 0.85)
    add_par(tf, text, 31, INK, bold=True, first=True)
    rrect(slide, 0.57, 1.14, 0.62, 0.09, fill=PINK, radius=0.045)


def chart_panel(slide, x, y, w, h, chart):
    rrect(slide, x, y, w, h, fill=SURFACE, line=LINEC, line_w=1.0, radius=0.16)
    a = ASPECT[chart]
    pad = 0.18
    cw = w - 2 * pad
    ch = cw / a
    if ch > h - 2 * pad:
        ch = h - 2 * pad
        cw = ch * a
    slide.shapes.add_picture(str(CHARTS / f"{chart}.png"),
                             Inches(x + (w - cw) / 2), Inches(y + (h - ch) / 2),
                             width=Inches(cw), height=Inches(ch))


nas = F["nasional"]
par = F["paradoks_pendidikan"]
muda = F["pengangguran_muda"]
kes = F["kesenjangan"]
prov = F["provinsi"]
kerja = F["pasar_kerja"]

# ---- Slide 1: Cover, Asymmetric Split + Data Hero ----
s = new_slide()
rrect(s, 7.62, 1.02, 5.35, 5.60, fill=PINKSOFT, radius=0.22)  # offset flat di belakang bezel
bezel(s, 7.38, 0.78, 5.35, 5.60)
tf = textbox(s, 7.68, 2.02, 4.75, 1.35)
add_par(tf, "17,37%", 68, PINK, bold=True, first=True, align=PP_ALIGN.CENTER)
tf = textbox(s, 7.68, 3.38, 4.75, 1.15)
add_par(tf, "TPT usia 15-24 tahun (Mei 2026)", 15, INK, bold=True,
        first=True, align=PP_ALIGN.CENTER)
add_par(tf, "Naik 1,01 poin persen dari Feb 2026, saat TPT nasional justru turun.",
        12.5, MUTED, align=PP_ALIGN.CENTER)
pill(s, 8.90, 5.42, 2.32, 0.54, "TPT nasional 4,65%", BLUE, WHITE, 12)

tf = textbox(s, 0.70, 1.02, 6.35, 1.95)
add_par(tf, "ANALISIS PENGANGGURAN INDONESIA 2026", 33, INK, bold=True, first=True)
tf = textbox(s, 0.70, 2.86, 6.20, 1.10)
add_par(tf, "Paradoks Pendidikan: TPT Nasional Turun, tapi Lulusan SMK Paling Sulit Dapat Kerja",
        17, BLUE, bold=True, first=True)
tf = textbox(s, 0.70, 4.06, 6.10, 1.45)
add_par(tf, "Data resmi BPS (Sakernas), Februari 2024 sampai Mei 2026. "
            "Dapat direproduksi dari script dan query SQL di repo.", 14, INK, first=True)
add_par(tf, "Plus: paket dashboard Power BI (star schema anti-jebakan, DAX, panduan rakit)",
        14, INK)
rrect(s, 0.72, 6.12, 0.62, 0.09, fill=PINK, radius=0.045)
tf = textbox(s, 0.70, 6.32, 6.35, 0.60)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 19, INK, bold=True, first=True)

# ---- Slide 2: Latar belakang & sumber data ----
s = new_slide()
title_block(s, "Latar Belakang dan Sumber Data")
bullets(s, 0.55, 1.52, 7.35, 5.20, [
    "Tingkat Pengangguran Terbuka (TPT) nasional turun, tapi angka rata-rata "
    "menyembunyikan jurang yang lebar antar jenjang pendidikan, umur, gender, "
    "kota-desa, dan provinsi.",
    "Project ini mengkompilasi angka resmi BPS dari publikasi Sakernas: "
    "5 titik waktu (Feb 2024, Feb 2025, Agu 2025, Feb 2026, Mei 2026), dengan "
    "rincian pendidikan, umur, gender, kota-desa, dan 38 provinsi.",
    "Semua angka diverifikasi dengan query SQL independen (sql/queries.sql) "
    "dan divisualisasikan dalam 10 chart dari data yang sama.",
    "Pertanyaan utama: TPT nasional 4,65%. Tapi siapa sebenarnya 7,22 juta "
    "penganggur itu, dan di mana letak masalah terbesarnya?",
])
bezel(s, 8.25, 1.52, 4.55, 4.42)
stats = [("5", "titik waktu Sakernas diamati"),
         ("38", "provinsi dalam kompilasi BPS"),
         ("10", "chart dari data yang sama")]
for i, (n, lab) in enumerate(stats):
    yy = 1.86 + i * 1.38
    tf = textbox(s, 8.62, yy, 1.25, 0.75)
    add_par(tf, n, 32, BLUE, bold=True, first=True)
    tf = textbox(s, 9.82, yy + 0.10, 2.80, 0.95)
    add_par(tf, lab, 13, INK, first=True)
    if i < 2:
        rrect(s, 8.62, yy + 1.02, 3.85, 0.022, fill=LINEC, radius=0.011)
tf = textbox(s, 0.55, 6.98, 12.2, 0.40)
add_par(tf, "Sumber: BPS, Survei Angkatan Kerja Nasional (Sakernas). "
            "Daftar sumber per angka: data/SOURCES.md.", 11, MUTED, italic=True, first=True)

# ---- Slide 3: TPT nasional ----
s = new_slide()
title_block(s, "TPT Nasional Turun, Ekonomi Tetap Tumbuh")
bullets(s, 0.55, 1.52, 5.85, 5.30, [
    f"TPT nasional turun dari {pct(nas['tpt_feb_2024'])} (Feb 2024) ke {pct(nas['tpt_mei_2026'])} "
    f"(Mei 2026), penurunan {num(abs(nas['penurunan_pp_feb24_mei26']))} poin persen dalam 27 bulan.",
    f"Penganggur Mei 2026: {jt(nas['penganggur_mei_2026_jt'])} orang dari angkatan kerja "
    f"{jt(nas['angkatan_kerja_mei_2026_jt'])}; yang bekerja {jt(nas['bekerja_mei_2026_jt'])}. "
    f"TPAK {pct(nas['tpak_mei_2026'])}.",
    f"Di saat yang sama, ekonomi tumbuh: PDB Q2 2026 tumbuh {pct(nas['pdb_q2_2026_yoy'])} "
    "secara tahunan. Penurunan TPT berjalan seiring pertumbuhan ekonomi.",
    "Tapi tren nasional saja tidak cukup, masalah sesungguhnya ada di rinciannya.",
])
chart_panel(s, 6.70, 1.48, 6.10, 3.90, "01_tpt_nasional_tren")
pill(s, 6.70, 5.62, 2.85, 0.55, "TPT Mei 2026: 4,65%", BLUE, WHITE, 12.5)
pill(s, 9.70, 5.62, 3.10, 0.55, "Penganggur: 7,22 jt orang", PINKSOFT, INK, 12.5)

# ---- Slide 4: Paradoks pendidikan ----
s = new_slide()
title_block(s, "Paradoks Pendidikan: Makin Tinggi Ijazah, Makin Sulit?")
bullets(s, 0.55, 1.52, 5.85, 5.30, [
    f"Lulusan SMK punya TPT {pct(par['smk_mei_2026'])}, tertinggi di antara semua "
    "jenjang pendidikan. Lulusan SD ke bawah hanya 2,31%.",
    f"Selisihnya {num(par['smk_vs_sd_pp'])} poin persen; TPT SMK {num(par['smk_vs_nasional_ratio'])} kali "
    "lipat rata-rata nasional. Paradoks: jenjang paling rendah pendidikannya "
    "justru paling mudah mendapat pekerjaan.",
    "Lulusan universitas (Diploma IV/S1/S2/S3) juga tinggi: 6,09%, lebih tinggi "
    "dari lulusan SMP. Pasar kerja tidak menyerap sesuai ijazah.",
    "Inilah fokus utama analisis ini: ketidakcocokan antara pendidikan dan "
    "lapangan kerja yang tersedia.",
])
rrect(s, 6.70, 1.48, 6.10, 1.32, fill=SURFACE, line=LINEC, line_w=1.0, radius=0.16)
tf = textbox(s, 6.98, 1.66, 2.20, 1.00)
add_par(tf, "7,68%", 40, PINK, bold=True, first=True)
tf = textbox(s, 9.20, 1.70, 3.40, 1.00)
add_par(tf, "TPT lulusan SMK (Mei 2026), tertinggi di antara semua jenjang pendidikan",
        13, INK, first=True)
chart_panel(s, 6.70, 2.98, 6.10, 3.95, "02_paradoks_pendidikan")

# ---- Slide 5: Tren per jenjang (chart di kiri) ----
s = new_slide()
title_block(s, "Tren SMK Membaik, tapi Masih Tertinggi")
chart_panel(s, 0.55, 1.48, 6.10, 3.90, "03_pendidikan_tren")
pill(s, 0.55, 5.62, 3.20, 0.55, "SMK: 8,62% ke 7,68%", PINK, WHITE, 12.5)
pill(s, 3.90, 5.62, 2.75, 0.55, "SD ke bawah: 2,31%", BLUE, WHITE, 12.5)
bullets(s, 6.95, 1.52, 5.85, 5.30, [
    f"TPT SMK turun dari {pct(par['smk_tren']['Feb 2024'])} (Feb 2024) ke "
    f"{pct(par['smk_mei_2026'])} (Mei 2026), turun {num(abs(par['smk_turun_pp_feb24_mei26']))} poin persen. "
    "Tren membaik.",
    "Tapi tetap jauh di atas semua jenjang lain: SMA 6,17%, universitas 6,09%, "
    "Diploma I/II/III 4,75%, SMP 4,16%, SD 2,31%.",
    "Lulusan universitas sempat naik dari 5,63% (Feb 2024) ke 6,23% (Feb 2025) "
    "sebelum turun ke 6,09%. Lulusan SD stabil di sekitar 2,3% sepanjang periode.",
    "Perbaikan tren belum menyentuh struktur masalahnya.",
])

# ---- Slide 6: Pengangguran muda (spotlight stat) ----
s = new_slide()
title_block(s, "Generasi Muda: 17,37% Menganggur")
bezel(s, 0.55, 1.48, 4.90, 2.52)
tf = textbox(s, 0.80, 1.66, 4.40, 1.05)
add_par(tf, "17,37%", 54, PINK, bold=True, first=True, align=PP_ALIGN.CENTER)
tf = textbox(s, 0.80, 2.76, 4.40, 0.45)
add_par(tf, "TPT usia 15-24 tahun (Mei 2026)", 13.5, INK, bold=True,
        first=True, align=PP_ALIGN.CENTER)
pill(s, 1.50, 3.30, 3.00, 0.48, "3,74x rata-rata nasional", PINKSOFT, INK, 12)
chart_panel(s, 0.55, 4.22, 4.90, 2.92, "04_umur")
bullets(s, 5.78, 1.52, 7.00, 5.30, [
    f"TPT usia 15-24 tahun: {pct(muda['muda_15_24_mei_2026'])} (Mei 2026), naik "
    f"{num(muda['muda_naik_pp_feb_mei_2026'])} poin persen dari Feb 2026, "
    "satu-satunya kelompok umur yang TPT-nya naik dalam periode itu.",
    f"Angka ini {num(muda['muda_vs_nasional_ratio'])} kali lipat rata-rata nasional dan "
    f"{num(muda['muda_vs_produktif_ratio'])} kali lipat usia produktif 25-59 tahun "
    f"({pct(muda['produktif_25_59_mei_2026'])}).",
    "Kelompok 25-59 tahun justru turun dari 3,08% (Feb 2024) ke 2,56%; kelompok "
    "60 tahun ke atas naik perlahan ke 2,35%.",
    "Krisis pengangguran Indonesia adalah krisis pengangguran muda, dan sebagian "
    "besar penganggur muda adalah lulusan SMA/SMK.",
])

# ---- Slide 7: Kesenjangan gender & kota-desa ----
s = new_slide()
title_block(s, "Kesenjangan Gender dan Kota-Desa")
bullets(s, 0.55, 1.52, 5.85, 5.30, [
    f"Pengangguran perkotaan {pct(kes['perkotaan_feb_2026'])} vs perdesaan "
    f"{pct(kes['perdesaan_feb_2026'])}, selisih {num(kes['urban_rural_gap_pp'])} poin persen, "
    f"kota {num(kes['urban_rural_ratio'])} kali lipat desa.",
    f"Laki-laki {pct(kes['laki_feb_2026'])} vs perempuan {pct(kes['perempuan_feb_2026'])} "
    f"(selisih {num(kes['gender_gap_pp'])} pp). Tapi TPAK perempuan hanya {pct(kes['tpak_perempuan_feb_2026'])} "
    f"vs laki-laki {pct(kes['tpak_laki_feb_2026'])}, banyak perempuan bahkan "
    "tidak masuk angkatan kerja.",
    "Artinya: pengangguran adalah persoalan perkotaan, dan di balik angka "
    "perempuan yang lebih rendah ada partisipasi yang jauh lebih kecil.",
])
bezel(s, 6.70, 1.48, 2.95, 1.90)
tf = textbox(s, 6.85, 1.72, 2.65, 0.80)
add_par(tf, "5,60%", 30, BLUE, bold=True, first=True, align=PP_ALIGN.CENTER)
tf = textbox(s, 6.85, 2.56, 2.65, 0.45)
add_par(tf, "Perkotaan (Feb 2026)", 12, INK, first=True, align=PP_ALIGN.CENTER)
bezel(s, 9.85, 1.48, 2.95, 1.90)
tf = textbox(s, 10.00, 1.72, 2.65, 0.80)
add_par(tf, "3,20%", 30, PINK, bold=True, first=True, align=PP_ALIGN.CENTER)
tf = textbox(s, 10.00, 2.56, 2.65, 0.45)
add_par(tf, "Perdesaan (Feb 2026)", 12, INK, first=True, align=PP_ALIGN.CENTER)
chart_panel(s, 6.70, 3.62, 6.10, 3.30, "05_kesenjangan")

# ---- Slide 8: Jurang provinsi ----
s = new_slide()
title_block(s, "Jurang Provinsi: Papua vs Bali 4,66 Kali Lipat")
bullets(s, 0.55, 1.52, 5.85, 4.60, [
    f"TPT tertinggi: Papua {pct(prov['top_mei_2026'][0]['tpt'])}; terendah: Bali "
    f"{pct(prov['terendah_mei_2026'][-1]['tpt'])}, selisih {num(prov['papua_vs_bali_ratio'])} kali lipat.",
    "Sepuluh provinsi tertinggi: " +
    ", ".join(f"{p['provinsi']} {pct(p['tpt'])}" for p in prov['top_mei_2026'][:3]) +
    "... Jawa Barat dan Banten ikut di papan atas meski pusat industri.",
    f"Penurunan terbesar Feb→Mei 2026: {prov['turun_terbesar_feb_mei_2026'][0]['provinsi']} "
    f"({num(prov['turun_terbesar_feb_mei_2026'][0]['perubahan_pp'])} pp). Kenaikan terbesar: "
    f"{prov['naik_terbesar_feb_mei_2026'][0]['provinsi']} "
    f"(+{num(prov['naik_terbesar_feb_mei_2026'][0]['perubahan_pp'])} pp).",
    "Catatan: Mei 2026 tersedia untuk 17 provinsi; Papua 2022-2023 tidak "
    "sebanding karena pemekaran.",
])
chart_panel(s, 6.70, 1.48, 6.10, 4.00, "06_provinsi_top10")
pill(s, 6.70, 5.72, 2.95, 0.55, "Tertinggi: Papua 7,08%", PINK, WHITE, 12.5)
pill(s, 9.85, 5.72, 2.95, 0.55, "Terendah: Bali 1,52%", BLUE, WHITE, 12.5)

# ---- Slide 9: Kualitas pekerjaan (dua chart) ----
s = new_slide()
title_block(s, "Bukan Sekadar Jumlah: Kualitas Pekerjaan")
tf = textbox(s, 0.55, 1.48, 6.00, 2.45)
add_bullet(tf, f"Pekerja formal hanya {pct(kerja['formal_feb_2026'])} ({jt(kerja['formal_jt_feb_2026'])}), "
               f"sisanya {pct(kerja['informal_feb_2026'])} bekerja informal.", size=15, first=True)
add_bullet(tf, f"Setengah menganggur {pct(kerja['setengah_menganggur_feb_2026'])}, pekerja paruh waktu "
               f"{pct(kerja['paruh_waktu_feb_2026'])}, komuter {jt(kerja['komuter_feb_2026_jt'])}.", size=15)
tf = textbox(s, 6.85, 1.48, 5.95, 2.45)
add_bullet(tf, "Upah rata-rata buruh naik dari Rp3,29 jt (Feb 2026) ke Rp3,39 jt (Mei 2026).",
           size=15, first=True)
add_bullet(tf, "Penambah lapangan kerja terbesar Feb→Mei 2026: " +
           ", ".join(f"{x['sektor'].lower()} (+{num(x['tambah_ribu'])} ribu)"
                     for x in kerja['sektor_penambah_feb_mei_2026']) + ".", size=15)
chart_panel(s, 0.55, 4.18, 6.90, 2.92, "09_formal_informal")
chart_panel(s, 7.70, 4.18, 5.10, 2.92, "10_upah_sektor")

# ---- Slide 10: Penutup, Designed Close ----
s = new_slide()
bezel(s, 0.55, 0.50, 12.23, 6.50)
tf = textbox(s, 1.00, 0.82, 11.30, 0.40)
add_par(tf, "PENUTUP", 12, MUTED, bold=True, first=True)
tf = textbox(s, 1.00, 1.22, 11.30, 1.35)
add_par(tf, "Kisah dua sisi 2026: TPT nasional turun ke 4,65% seiring ekonomi yang "
            "tumbuh 5,29%, tapi lulusan SMK (7,68%) dan anak muda 15-24 tahun (17,37%) "
            "tetap jadi kelompok paling tertekan.", 17.5, INK, first=True)
pill(s, 1.00, 2.62, 2.85, 0.55, "TPT nasional 4,65%", BLUE, WHITE, 12.5)
pill(s, 4.00, 2.62, 2.30, 0.55, "SMK 7,68%", PINK, WHITE, 12.5)
pill(s, 6.45, 2.62, 3.05, 0.55, "Muda 15-24: 17,37%", PINKSOFT, INK, 12.5)
tf = textbox(s, 1.00, 3.42, 11.30, 1.55)
add_par(tf, "Isi repo: analysis.ipynb, data/compiled + SOURCES.md, sql/queries.sql "
            "(semua angka diverifikasi ulang lewat SQL), scripts, charts, findings.json.",
        13.5, MUTED, first=True)
add_par(tf, "Plus kit dashboard Power BI: star schema anti-jebakan (fact_tpt_nasional "
            "dan fact_tpt_segmen terpisah agar measure tidak terkontaminasi), measure DAX, "
            "layout dashboard, dan panduan rakit 20-30 menit. Lihat "
            "powerbi/measures_dax_CATATAN_LEVEL.md dulu.", 13.5, MUTED)
tf = textbox(s, 1.00, 5.02, 11.30, 0.45)
add_par(tf, "Keterbatasan: angka Mei 2026 hanya untuk 17 provinsi; dataset ini kompilasi "
            "angka publikasi resmi, bukan mikrodataset Sakernas.", 12.5, MUTED, italic=True, first=True)
rrect(s, 1.02, 5.62, 0.62, 0.09, fill=PINK, radius=0.045)
tf = textbox(s, 1.00, 5.82, 11.30, 0.70)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 27, INK, bold=True, first=True)
tf = textbox(s, 1.00, 6.52, 11.30, 0.45)
p = tf.paragraphs[0]
p.space_after = Pt(0)
r = p.add_run()
r.text = "github.com/baghizash/pengangguran-indonesia-2026"
r.font.size = Pt(14)
r.font.bold = True
r.font.name = FONT
r.font.color.rgb = BLUE
r = p.add_run()
r.text = "  |  Terbuka untuk peluang data analyst dan kolaborasi analisis data publik."
r.font.size = Pt(14)
r.font.name = FONT
r.font.color.rgb = INK

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
shutil.copy(OUT, OUT_COPY)
print(f"PPT saved: {OUT}")
print(f"Copy saved: {OUT_COPY}")
